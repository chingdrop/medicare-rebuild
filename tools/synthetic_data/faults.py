"""Named faults for proving that reconciliation catches real problems.

Faults are applied AFTER a demo run, to the loaded database or to the report file.
They never change the pipeline, the generated source files, or the manifest, and
they are off unless the generator was asked for them (`--inject-fault`).
"""

from pathlib import Path

from openpyxl import load_workbook

FAULTS = {
    "drop-rows": "delete three glucose readings of a patient with no billed codes",
    "duplicate-keys": "give two devices the same hardware ID",
    "orphan-fk": "point one blood pressure reading at a device that does not exist",
    "report-off-by-one": "add one to a code count in the billing report file",
}

# Which reconciliation check each fault is designed to trip.
EXPECTED_CHECK = {
    "drop-rows": "1 row conservation",
    "duplicate-keys": "2a key uniqueness",
    "orphan-fk": "2b foreign keys",
    "report-off-by-one": "6 report totals",
}

# Each fault is a list of one or more statements, run in order.
_SQL: dict[str, list[str]] = {
    "drop-rows": [
        """
        DELETE FROM glucose_reading WHERE glucose_reading_id IN (
            SELECT TOP 3 gr.glucose_reading_id
            FROM glucose_reading gr JOIN device d ON d.device_id = gr.device_id
            WHERE d.patient_id NOT IN (SELECT patient_id FROM medical_code)
            ORDER BY d.patient_id, gr.glucose_reading_id)"""
    ],
    "duplicate-keys": [
        """
        UPDATE device
        SET hardware_uuid = (
            SELECT hardware_uuid FROM device
            WHERE device_id = (SELECT MIN(device_id) FROM device))
        WHERE device_id = (SELECT MAX(device_id) FROM device)"""
    ],
    "orphan-fk": [
        # models.py declares a real foreign key on device_id (decision 0015), unlike
        # the demo's earlier hand-written schema, so it has to be disabled before this
        # deliberately-corrupting UPDATE can run -- exactly what a real DBA would need
        # to do to hand-craft a broken row like this one.
        "ALTER TABLE blood_pressure_reading NOCHECK CONSTRAINT ALL",
        """
        UPDATE blood_pressure_reading SET device_id = 987654321
        WHERE blood_pressure_reading_id = (
            SELECT TOP 1 b.blood_pressure_reading_id
            FROM blood_pressure_reading b JOIN device d ON d.device_id = b.device_id
            WHERE d.patient_id NOT IN (SELECT patient_id FROM medical_code)
            ORDER BY b.blood_pressure_reading_id)""",
    ],
}


def bump_report_count(report_path: Path) -> None:
    """Add one to the first positive 99457 (else any code) count in the report."""
    wb = load_workbook(report_path)
    ws = wb.active
    header = {str(c.value): c.column for c in ws[1]}
    for code in ("99457", "99454", "99453", "99202", "99458"):
        col = header.get(code)
        if col is None:
            continue
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            if isinstance(cell.value, int | float) and cell.value > 0:
                cell.value = int(cell.value) + 1
                wb.save(report_path)
                return
    raise RuntimeError("no positive code count found to alter")


def apply_faults(names: list[str], report_path: Path) -> None:
    from tools.synthetic_data.demo import GPS_DB, _run

    for name in names:
        if name not in FAULTS:
            raise ValueError(f"unknown fault {name!r}; choose from {sorted(FAULTS)}")
        if name == "report-off-by-one":
            bump_report_count(report_path)
        else:
            _run(GPS_DB, _SQL[name])
