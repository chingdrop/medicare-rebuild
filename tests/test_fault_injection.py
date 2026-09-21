"""Unit tests for fault injection (the generator side). No database needed."""

import json

import pytest
from openpyxl import Workbook, load_workbook

from tools.synthetic_data import config as cfg
from tools.synthetic_data.faults import FAULTS, bump_report_count
from tools.synthetic_data.generator import FAULTS_FILE, generate, minimum_patients


def test_faults_are_off_by_default(tmp_path):
    generate(cfg.DEFAULT_SEED, minimum_patients(), tmp_path)
    assert not (tmp_path / FAULTS_FILE).exists()


def test_a_requested_fault_is_recorded_and_stale_ones_are_cleared(tmp_path):
    generate(cfg.DEFAULT_SEED, minimum_patients(), tmp_path, ["orphan-fk"])
    assert json.loads((tmp_path / FAULTS_FILE).read_text()) == {"faults": ["orphan-fk"]}
    generate(cfg.DEFAULT_SEED, minimum_patients(), tmp_path)
    assert not (tmp_path / FAULTS_FILE).exists()


def test_an_unknown_fault_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="unknown fault"):
        generate(cfg.DEFAULT_SEED, minimum_patients(), tmp_path, ["not-a-fault"])


def test_injecting_a_fault_does_not_change_the_generated_source_files(tmp_path):
    clean, faulty = tmp_path / "clean", tmp_path / "faulty"
    generate(cfg.DEFAULT_SEED, minimum_patients(), clean)
    generate(cfg.DEFAULT_SEED, minimum_patients(), faulty, sorted(FAULTS))
    for name in ("Patient_Export.csv", "manifest.json", "legacy/Glucose_Readings.csv"):
        assert (clean / name).read_bytes() == (faulty / name).read_bytes()


def test_report_off_by_one_changes_exactly_one_count_by_one(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "99202", "99453", "99454", "99457", "99458"])
    ws.append([1, 0, 1, 1, 0, 0])
    ws.append([2, 0, 0, 0, 2, 0])
    path = tmp_path / "Billing_Report.xlsx"
    wb.save(path)
    bump_report_count(path)
    rows = [
        [c.value for c in r] for r in load_workbook(path).active.iter_rows(min_row=2)
    ]
    assert rows == [[1, 0, 1, 1, 0, 0], [2, 0, 0, 0, 3, 0]]
